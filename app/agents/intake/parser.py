"""Resume PDF/DOCX text extraction + LLM structured extraction.

Phase 1 keeps this independent of app/utils via an injected `generate_structured`
callable, wired in once app/utils/llm_client.py exists.
"""

from enum import Enum
from io import BytesIO
from pathlib import Path
from typing import Callable, Awaitable

import fitz
from docx import Document
from pydantic import BaseModel, Field

from app.core.schemas import ProficiencyLevel


class ResumeParseError(ValueError):
    """Raised when a resume cannot be parsed."""


class ResumeSkill(BaseModel):
    skill_name: str
    category: str
    proficiency: ProficiencyLevel
    evidence: str | None = None


class ResumeProfile(BaseModel):
    """Intake-internal — not part of app/core/schemas.py (never crosses to Module 2)."""
    name: str | None = None
    summary: str | None = None
    skills: list[ResumeSkill] = Field(default_factory=list)
    experience: list[str] = Field(default_factory=list)
    projects: list[str] = Field(default_factory=list)
    education: list[str] = Field(default_factory=list)


def extract_resume_text(filename: str, content: bytes) -> str:
    """Extract plain text from a PDF or DOCX resume."""
    extension = Path(filename).suffix.lower()

    if extension == ".pdf":
        try:
            with fitz.open(stream=content, filetype="pdf") as document:
                text = "\n".join(page.get_text("text") for page in document)
        except Exception as exc:
            raise ResumeParseError("Unable to parse PDF resume.") from exc

    elif extension == ".docx":
        try:
            document = Document(BytesIO(content))
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            table_rows = [
                " | ".join(cell.text.strip() for cell in row.cells)
                for table in document.tables
                for row in table.rows
            ]
            text = "\n".join(paragraphs + table_rows)
        except Exception as exc:
            raise ResumeParseError("Unable to parse DOCX resume.") from exc

    else:
        raise ResumeParseError("Only PDF and DOCX resumes are supported.")

    text = text.strip()

    if not text:
        raise ResumeParseError("Resume contains no extractable text.")

    return text


async def parse_resume(
    filename: str,
    content: bytes,
    generate_structured: Callable[..., Awaitable] | None = None,
) -> tuple[str, ResumeProfile]:
    """Extract resume text, then use the injected LLM callable for structuring."""
    resume_text = extract_resume_text(filename, content)

    if generate_structured is None:
        raise RuntimeError(
            "No LLM callable supplied. Connect the project's shared LLM client "
            "when app/utils/llm_client.py is created."
        )

    from app.agents.intake.prompts import RESUME_EXTRACTION_SYSTEM

    result = await generate_structured(
        system_prompt=RESUME_EXTRACTION_SYSTEM,
        user_prompt=f"RESUME TEXT:\n{resume_text}",
        response_model=ResumeProfile,
    )

    profile = result if isinstance(result, ResumeProfile) else ResumeProfile.model_validate(result)
    return resume_text, profile